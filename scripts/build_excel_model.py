

import os
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DATA_DIR = "data/processed"
OUT_PATH = "reports/churn_intervention_model.xlsx"

# Business assumptions (you choose these — be ready to defend them)
CAMPAIGN_COST_PER_CUSTOMER = 15.0   # £ cost to contact one at-risk customer
BASE_SUCCESS_RATE          = 0.25   # 25% of contacted customers return
DISCOUNT_OFFERED           = 0.10   # 10% off to re-engage them
CONSERVATIVE_SUCCESS       = 0.15
OPTIMISTIC_SUCCESS         = 0.35



def fill(hex_colour):
    return PatternFill("solid", fgColor=hex_colour)

def thin_border():
    s = Side(style="thin", color="D0D7DE")
    return Border(left=s, right=s, top=s, bottom=s)

def title_bar(ws, cell_range, text):
    ws.merge_cells(cell_range)
    c = ws[cell_range.split(":")[0]]
    c.value = text
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = fill("0F172A")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[c.row].height = 26

def header_row(ws, row, headers, start_col=1, colour="1E293B"):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = fill(colour)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border()
    ws.row_dimensions[row].height = 24

def put(ws, row, col, value, fmt=None, bold=False, bg=None, align="center"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=10)
    c.border = thin_border()
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
    if fmt:
        c.number_format = fmt
    if bg:
        c.fill = fill(bg)
    return c

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w



print("Loading data...")

churn = pd.read_csv(f"{DATA_DIR}/churn_scores.csv")
feats = pd.read_csv(f"{DATA_DIR}/customer_features.csv")

print(f"  churn_scores      : {churn.shape}  cols={churn.columns.tolist()}")
print(f"  customer_features : {feats.shape}")

# --- find the risk band column ---
risk_col = None
for candidate in ["risk_band", "riskband", "risk", "band"]:
    if candidate in churn.columns:
        risk_col = candidate
        break
if risk_col is None:
    raise SystemExit(
        "Could not find a risk band column in churn_scores.csv.\n"
        f"Columns present: {churn.columns.tolist()}\n"
        "Re-run Notebook 04 Cell 10, or rename your column to 'risk_band'."
    )

# --- find the churn probability column ---
prob_col = None
for candidate in ["churn_probability", "churn_proba", "probability", "churn_prob"]:
    if candidate in churn.columns:
        prob_col = candidate
        break

# --- find the monetary / spend column ---
money_col = None
for candidate in ["monetary", "total_spend", "spend", "revenue"]:
    if candidate in feats.columns:
        money_col = candidate
        break
if money_col is None:
    raise SystemExit(
        "Could not find a spend column in customer_features.csv.\n"
        f"Columns present: {feats.columns.tolist()}"
    )

# --- merge so every at-risk customer has a spend figure attached ---
merged = churn.merge(
    feats[["customer_id", money_col]],
    on="customer_id",
    how="left"
)
merged[money_col] = merged[money_col].fillna(0)

total_customers = len(merged)

high  = merged[merged[risk_col].astype(str).str.lower() == "high"]
med   = merged[merged[risk_col].astype(str).str.lower() == "medium"]
low   = merged[merged[risk_col].astype(str).str.lower() == "low"]

high_n   = len(high)
high_rev = float(high[money_col].sum())
high_avg = float(high[money_col].mean()) if high_n else 0.0

total_rev = float(merged[money_col].sum())

print(f"\n  Total customers scored : {total_customers:,}")
print(f"  High-risk customers    : {high_n:,}  ({high_n/total_customers:.1%})")
print(f"  Revenue at risk        : {high_rev:,.0f}")
print(f"  Avg spend, at-risk     : {high_avg:,.0f}")

if high_n == 0:
    print("\n  WARNING: zero high-risk customers found.")
    print(f"  Values present in '{risk_col}': {merged[risk_col].unique().tolist()}")
    print("  The workbook will still build but the model will be empty.")


wb = Workbook()

ws = wb.active
ws.title = "Assumptions"
set_widths(ws, [42, 18, 52])

title_bar(ws, "A1:C1", "MODEL ASSUMPTIONS  —  edit the blue cells, everything recalculates")

header_row(ws, 3, ["Parameter", "Value", "Source / Justification"])

assumptions = [
    ("Total customers scored",        total_customers, "0",       "F8FAFC",
     "Output of the Logistic Regression churn model"),
    ("High-risk customers",           high_n,          "0",       "F8FAFC",
     "Predicted churn probability > 0.70"),
    ("Avg annual spend, at-risk seg.", high_avg,       "#,##0.00","F8FAFC",
     "Mean historical spend of the high-risk cohort"),
    ("Total revenue at risk",         high_rev,        "#,##0.00","F8FAFC",
     "Sum of historical spend across high-risk customers"),
    ("Campaign cost per customer",    CAMPAIGN_COST_PER_CUSTOMER, "#,##0.00", "DBEAFE",
     "ASSUMPTION: email + offer fulfilment cost"),
    ("Win-back success rate (base)",  BASE_SUCCESS_RATE, "0.0%",  "DBEAFE",
     "ASSUMPTION: industry benchmark for retail win-back"),
    ("Discount offered to re-engage", DISCOUNT_OFFERED,  "0.0%",  "DBEAFE",
     "ASSUMPTION: margin given up to trigger return"),
    ("Conservative success rate",     CONSERVATIVE_SUCCESS, "0.0%","DBEAFE",
     "ASSUMPTION: downside scenario"),
    ("Optimistic success rate",       OPTIMISTIC_SUCCESS,   "0.0%","DBEAFE",
     "ASSUMPTION: upside scenario"),
]

for i, (label, val, fmt, bg, note) in enumerate(assumptions, start=4):
    put(ws, i, 1, label, bold=True, bg=bg, align="left")
    put(ws, i, 2, val, fmt=fmt, bold=True, bg=bg)
    put(ws, i, 3, note, bg=bg, align="left")
    ws.row_dimensions[i].height = 18

# Legend
ws.merge_cells("A14:C14")
c = ws["A14"]
c.value = ("Grey rows are computed from the model output. Blue rows are business assumptions "
           "chosen by the analyst — change them and every sheet updates.")
c.font = Font(italic=True, size=9, color="475569")
c.fill = fill("FEF3C7")
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[14].height = 28

# Named cell references used by later sheets
A_HIGH_N   = "Assumptions!$B$5"
A_AVG_SPD  = "Assumptions!$B$6"
A_COST     = "Assumptions!$B$8"
A_BASE     = "Assumptions!$B$9"
A_DISC     = "Assumptions!$B$10"
A_CONS     = "Assumptions!$B$11"
A_OPT      = "Assumptions!$B$12"



ws2 = wb.create_sheet("Intervention ROI")
set_widths(ws2, [40, 20, 20, 20, 34])

title_bar(ws2, "A1:E1", "WIN-BACK CAMPAIGN  —  ROI MODEL")

ws2.merge_cells("A2:E2")
c = ws2["A2"]
c.value = ("Question: is it worth running a win-back campaign against the high-risk cohort, "
           "and what success rate does it need to break even?")
c.font = Font(italic=True, size=9, color="374151")
c.fill = fill("F8FAFC")
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws2.row_dimensions[2].height = 22

header_row(ws2, 4, ["Metric", "Conservative", "Base Case", "Optimistic", "Notes"])

# Column letters for the three scenarios
SC = {"B": A_CONS, "C": A_BASE, "D": A_OPT}

rows = [
    # (label, formula_template, number_format, background, note)
    ("Customers targeted",
     f"={A_HIGH_N}", "#,##0", "F8FAFC",
     "Entire high-risk cohort"),

    ("Assumed success rate",
     None, "0.0%", "DBEAFE",
     "Share of contacted customers who return"),

    ("Expected customers recovered",
     "={c}5*{c}6", "#,##0", "F8FAFC",
     "Targeted x success rate"),

    ("Gross revenue recovered",
     f"={{c}}7*{A_AVG_SPD}", "#,##0", "F8FAFC",
     "Recovered customers x avg annual spend"),

    ("Less: discount given",
     f"=-{{c}}8*{A_DISC}", "#,##0", "FEF3C7",
     "Margin sacrificed to trigger return"),

    ("Less: campaign cost",
     f"=-{{c}}5*{A_COST}", "#,##0", "FEF3C7",
     "Cost of contacting every targeted customer"),

    ("Net benefit",
     "={c}8+{c}9+{c}10", "#,##0", "DCFCE7",
     "Revenue recovered minus all costs"),

    ("Total campaign cost",
     f"={{c}}5*{A_COST}", "#,##0", "F8FAFC",
     "Denominator for ROI"),

    ("ROI",
     "={c}11/{c}12", "0.0%", "DCFCE7",
     "Net benefit / campaign spend"),

    ("Break-even success rate",
     f"={A_COST}/({A_AVG_SPD}*(1-{A_DISC}))", "0.00%", "EDE9FE",
     "Minimum success rate for the campaign to pay for itself"),
]

for r_offset, (label, formula, fmt, bg, note) in enumerate(rows):
    r = 5 + r_offset
    put(ws2, r, 1, label, bold=True, bg=bg, align="left")

    for col_letter, scenario_ref in SC.items():
        col_idx = {"B": 2, "C": 3, "D": 4}[col_letter]
        if label == "Assumed success rate":
            value = f"={scenario_ref}"
        elif formula is None:
            value = ""
        else:
            value = formula.replace("{c}", col_letter)
        put(ws2, r, col_idx, value, fmt=fmt, bg=bg,
            bold=(label in ("Net benefit", "ROI")))

    put(ws2, r, 5, note, bg=bg, align="left")
    ws2.row_dimensions[r].height = 18

# Interpretation box
ws2.merge_cells("A16:E17")
c = ws2["A16"]
c.value = ("HOW TO READ THIS: The break-even success rate is the single most useful number here. "
           "If the campaign can realistically beat it, the campaign is worth running. If it cannot, "
           "the correct decision is not to run it — the model's value is telling you that before "
           "you spend the budget, not after.")
c.font = Font(italic=True, size=9, color="374151")
c.fill = fill("FEF3C7")
c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)



ws3 = wb.create_sheet("Sensitivity")
set_widths(ws3, [26, 16, 16, 16, 16, 16, 16, 16])

title_bar(ws3, "A1:H1", "SENSITIVITY  —  NET BENEFIT ACROSS SUCCESS RATES")

ws3.merge_cells("A2:H2")
c = ws3["A2"]
c.value = ("The base case assumes one success rate. This grid shows what happens across the "
           "plausible range, so the decision isn't hostage to a single assumption.")
c.font = Font(italic=True, size=9, color="374151")
c.fill = fill("F8FAFC")
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws3.row_dimensions[2].height = 22

rates = [0.05, 0.10, 0.15, 0.20, 0.25, 0.30, 0.40]

header_row(ws3, 4, ["Success rate →"] + [f"{r:.0%}" for r in rates])

# Row of the rate values themselves (used by formulas below)
put(ws3, 5, 1, "Rate", bold=True, bg="F1F5F9", align="left")
for i, r in enumerate(rates):
    put(ws3, 5, 2 + i, r, fmt="0%", bg="F1F5F9")

sens_rows = [
    ("Customers recovered",
     f"={A_HIGH_N}*{{c}}5", "#,##0", "F8FAFC"),
    ("Gross revenue recovered",
     f"={{c}}6*{A_AVG_SPD}", "#,##0", "F8FAFC"),
    ("Net benefit",
     f"={{c}}7*(1-{A_DISC})-{A_HIGH_N}*{A_COST}", "#,##0", "DCFCE7"),
    ("ROI",
     f"={{c}}8/({A_HIGH_N}*{A_COST})", "0.0%", "DCFCE7"),
]

for r_offset, (label, formula, fmt, bg) in enumerate(sens_rows):
    r = 6 + r_offset
    put(ws3, r, 1, label, bold=True, bg=bg, align="left")
    for i in range(len(rates)):
        col_letter = get_column_letter(2 + i)
        put(ws3, r, 2 + i, formula.replace("{c}", col_letter), fmt=fmt, bg=bg)
    ws3.row_dimensions[r].height = 18

ws3.merge_cells("A11:H12")
c = ws3["A11"]
c.value = ("Read across the 'Net benefit' row to find where it crosses zero. That crossing point "
           "is the break-even success rate — everything to the right is profitable, everything "
           "to the left destroys value.")
c.font = Font(italic=True, size=9, color="374151")
c.fill = fill("FEF3C7")
c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)



ws4 = wb.create_sheet("Risk Segments")
set_widths(ws4, [20, 16, 16, 20, 20, 34])

title_bar(ws4, "A1:F1", "CUSTOMER BASE BY CHURN RISK BAND")

header_row(ws4, 3, ["Risk Band", "Customers", "% of Base",
                    "Revenue Represented", "Avg Spend", "Recommended Action"])

band_actions = {
    "High":   ("Targeted win-back campaign — highest urgency", "FEE2E2"),
    "Medium": ("Monitor; low-cost re-engagement (newsletter, offers)", "FEF3C7"),
    "Low":    ("No intervention needed — protect via normal service", "DCFCE7"),
}

for i, (band, subset) in enumerate([("High", high), ("Medium", med), ("Low", low)], start=4):
    n = len(subset)
    rev = float(subset[money_col].sum())
    avg = float(subset[money_col].mean()) if n else 0.0
    action, bg = band_actions[band]

    put(ws4, i, 1, band, bold=True, bg=bg, align="left")
    put(ws4, i, 2, n, fmt="#,##0", bg=bg)
    put(ws4, i, 3, (n / total_customers) if total_customers else 0, fmt="0.0%", bg=bg)
    put(ws4, i, 4, rev, fmt="#,##0", bg=bg)
    put(ws4, i, 5, avg, fmt="#,##0", bg=bg)
    put(ws4, i, 6, action, bg=bg, align="left")
    ws4.row_dimensions[i].height = 18

# Totals row
put(ws4, 7, 1, "TOTAL", bold=True, bg="1E293B", align="left").font = Font(bold=True, size=10, color="FFFFFF")
put(ws4, 7, 2, total_customers, fmt="#,##0", bold=True, bg="E2E8F0")
put(ws4, 7, 3, 1.0, fmt="0.0%", bold=True, bg="E2E8F0")
put(ws4, 7, 4, total_rev, fmt="#,##0", bold=True, bg="E2E8F0")
put(ws4, 7, 5, (total_rev / total_customers) if total_customers else 0, fmt="#,##0", bold=True, bg="E2E8F0")
put(ws4, 7, 6, "", bg="E2E8F0")



os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
wb.save(OUT_PATH)

print(f"\n  Saved: {OUT_PATH}")
print(f"  Sheets: {[s.title for s in wb.worksheets]}")
print("\nOpen it in Excel and check the Intervention ROI sheet calculates.")